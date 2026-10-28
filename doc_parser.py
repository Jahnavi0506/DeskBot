"""Document and Image Parser for DRL Coding Agent.

Extracts questions and text content from:
- Images (.png, .jpg, .jpeg, .bmp, .webp) using Windows built-in OCR (winsdk.windows.media.ocr)
- PDFs (.pdf) using pypdf
- Word documents (.docx) using python-docx
- Plain text / source code (.txt, .py, .csv, .json, .md, etc.)
"""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from safety import resolve_user_path


def extract_text_from_file(file_path: str | Path) -> str:
    """Extract readable text from an image, PDF, DOCX, or text file."""
    path = resolve_user_path(file_path) if isinstance(file_path, str) else file_path.resolve()
    if not path.exists() or not path.is_file():
        return f"File not found: {path}"

    suffix = path.suffix.lower()

    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
        return extract_text_from_image(path)
    elif suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix in {".docx", ".doc"}:
        return extract_text_from_docx(path)
    else:
        # Fallback to plain text
        try:
            return path.read_text(encoding="utf-8", errors="replace").strip()
        except Exception as exc:
            return f"Error reading text file: {exc}"


def extract_text_from_image(image_path: Path) -> str:
    """Run Windows Media OCR on an image file to extract textual questions."""
    try:
        import winsdk.windows.graphics.imaging as imaging
        import winsdk.windows.media.ocr as ocr
        import winsdk.windows.storage.streams as streams

        async def _ocr_coro() -> str:
            bytes_data = image_path.read_bytes()
            writer = streams.DataWriter()
            writer.write_bytes(bytes_data)
            ibuffer = writer.detach_buffer()

            stream = streams.InMemoryRandomAccessStream()
            await stream.write_async(ibuffer)
            stream.seek(0)

            decoder = await imaging.BitmapDecoder.create_async(stream)
            bitmap = await decoder.get_software_bitmap_async()

            engine = ocr.OcrEngine.try_create_from_user_profile_languages()
            if not engine:
                import winsdk.windows.globalization as glob
                engine = ocr.OcrEngine.try_create_from_language(glob.Language("en-US"))
            if not engine:
                return "OCR Engine unavailable on this system."

            result = await engine.recognize_async(bitmap)
            return (result.text or "").strip()

        # Safely run in sync context or inside running loop
        try:
            loop = asyncio.get_event_loop()
            if loop.is_running():
                import concurrent.futures
                with concurrent.futures.ThreadPoolExecutor() as pool:
                    return pool.submit(asyncio.run, _ocr_coro()).result(timeout=25)
            else:
                return loop.run_until_complete(_ocr_coro())
        except Exception:
            return asyncio.run(_ocr_coro())
    except Exception as exc:
        return f"[Image OCR Error on {image_path.name}: {exc}]"


def extract_text_from_pdf(pdf_path: Path) -> str:
    """Extract text from PDF pages using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(str(pdf_path))
        text_parts: list[str] = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(f"--- Page {i+1} ---\n{page_text.strip()}")
        return "\n\n".join(text_parts).strip() or "(No extractable text found in PDF)"
    except Exception as exc:
        return f"[PDF Extraction Error on {pdf_path.name}: {exc}]"


def extract_text_from_docx(docx_path: Path) -> str:
    """Extract paragraphs and tables from a Word (.docx) document."""
    try:
        import docx
        doc = docx.Document(str(docx_path))
        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for t in doc.tables:
            for row in t.rows:
                row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_vals:
                    parts.append(" | ".join(row_vals))
        return "\n".join(parts).strip() or "(Empty document)"
    except Exception as exc:
        return f"[DOCX Extraction Error on {docx_path.name}: {exc}]"
